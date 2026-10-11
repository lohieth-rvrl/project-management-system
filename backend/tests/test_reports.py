import io

import pytest
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import User
from apps.projects.models import Project
from apps.risks.models import Risk
from apps.tasks.models import Task

pytestmark = pytest.mark.django_db


@pytest.fixture
def proj():
    u = User.objects.create_user("rep", password="x" * 10, role="viewer")
    p = Project.objects.create(name="Tom & Jerry <1>", code="RP", budget=1000, hourly_rate=50)
    Task.objects.create(project=p, title="Write <b>docs</b> & ship", status="done", estimate_hours=3)
    Task.objects.create(project=p, title="Second", status="todo")
    Risk.objects.create(project=p, title="Vendor slip", probability=4, impact=4)
    c = APIClient()
    c.force_authenticate(u)
    return c, p


def test_xlsx_report_contents(proj):
    c, p = proj
    r = c.get(f"/api/analytics/projects/{p.id}/report/?file=xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r["Content-Type"]
    assert 'RP-report.xlsx' in r["Content-Disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Tasks", "Time by person", "Open risks"]
    assert [row[1].value for row in wb["Tasks"].iter_rows(min_row=2)].count("Second") == 1
    assert wb["Open risks"]["A2"].value == "Vendor slip"


def test_pdf_report_handles_special_characters(proj):
    c, p = proj
    r = c.get(f"/api/analytics/projects/{p.id}/report/?file=pdf")
    assert r.status_code == 200 and r.content.startswith(b"%PDF")


def test_empty_project_and_bad_format(proj):
    c, _ = proj
    e = Project.objects.create(name="E", code="EM")
    assert c.get(f"/api/analytics/projects/{e.id}/report/?file=pdf").status_code == 200
    assert c.get(f"/api/analytics/projects/{e.id}/report/?file=xlsx").status_code == 200
    assert c.get(f"/api/analytics/projects/{e.id}/report/?file=doc").status_code == 400
    assert APIClient().get(f"/api/analytics/projects/{e.id}/report/").status_code == 401
