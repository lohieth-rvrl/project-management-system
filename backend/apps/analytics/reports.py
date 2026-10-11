"""Project status report as an Excel workbook or a PDF."""
import io

from apps.timesheets.models import TimeEntry

STATUS = {"todo": "To Do", "in_progress": "In Progress", "review": "In Review", "blocked": "Blocked", "done": "Done"}


def _rows(project):
    tasks = project.tasks.select_related("assignee").order_by("status", "due_date", "id")
    return [[t.id, t.title, STATUS.get(t.status, t.status), t.get_priority_display(),
             t.assignee.username if t.assignee else "", t.start_date, t.due_date, float(t.estimate_hours)]
            for t in tasks]


def _logged(project):
    out = {}
    for e in TimeEntry.objects.filter(task__project=project).select_related("user"):
        out[e.user.username] = out.get(e.user.username, 0) + float(e.hours)
    return sorted(out.items())


def _kpis(project, s):
    pct = f'{s["budget_used_percent"]}%' if s["budget_used_percent"] is not None else "n/a"
    return [
        ("Status", project.get_status_display()), ("Health", project.get_health_display()),
        ("Progress", f'{s["progress_percent"]}% ({s["done_tasks"]}/{s["total_tasks"]} tasks)'),
        ("Overdue tasks", s["overdue_tasks"]), ("Estimated hours", s["estimated_hours"]),
        ("Logged hours", s["logged_hours"]), ("Budget", s["budget"]),
        ("Cost to date", s["cost_to_date"]), ("Budget used", pct),
        ("Open risks", s["open_risks"]), ("Missed or late milestones", s["milestones_missed"]),
    ]


def _risks(project):
    return [[r.title, r.score, r.get_status_display() if hasattr(r, "get_status_display") else r.status]
            for r in project.risks.exclude(status="closed").order_by("-probability", "-impact")]


def build_xlsx(project, s):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F4E78")

    def table(ws, header, rows, start=1):
        for c, h in enumerate(header, 1):
            cell = ws.cell(row=start, column=c, value=h)
            cell.font, cell.fill = head, fill
        for r, row in enumerate(rows, start + 1):
            for c, v in enumerate(row, 1):
                ws.cell(row=r, column=c, value=v)
        for c in range(1, len(header) + 1):
            width = max(len(str(ws.cell(row=r, column=c).value or "")) for r in range(start, start + len(rows) + 1))
            ws.column_dimensions[get_column_letter(c)].width = min(max(12, width + 2), 60)

    ws = wb.active
    ws.title = "Summary"
    ws["A1"], ws["A1"].font = f"{project.code} - {project.name}", Font(bold=True, size=14)
    table(ws, ["Measure", "Value"], _kpis(project, s), start=3)
    table(wb.create_sheet("Tasks"), ["ID", "Title", "Status", "Priority", "Assignee", "Start", "Due", "Estimate (h)"],
          _rows(project))
    table(wb.create_sheet("Time by person"), ["Person", "Hours"], _logged(project))
    table(wb.create_sheet("Open risks"), ["Risk", "Score", "Status"], _risks(project))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(project, s):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    cell = styles["BodyText"]
    cell.fontSize = 8
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), title=f"{project.code} report",
                            leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30)

    def tbl(data, widths=None):
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                               ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTSIZE", (0, 0), (-1, -1), 8),
                               ("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        return t

    esc = lambda v: str(v if v is not None else "").replace("&", "&amp;").replace("<", "&lt;")  # noqa: E731
    story = [Paragraph(esc(f"{project.code} - {project.name}"), styles["Title"]),
             tbl([["Measure", "Value"]] + [[a, esc(b)] for a, b in _kpis(project, s)], [200, 300]),
             Spacer(1, 14), Paragraph("Tasks", styles["Heading2"])]
    rows = [[esc(c) for c in r] for r in _rows(project)]
    story.append(tbl([["ID", "Title", "Status", "Priority", "Assignee", "Start", "Due", "Est (h)"]] +
                     [[r[0], Paragraph(r[1], cell)] + r[2:] for r in rows] if rows else [["No tasks"]],
                     [35, 300, 70, 60, 80, 65, 65, 50] if rows else None))
    risks = _risks(project)
    if risks:
        story += [Spacer(1, 14), Paragraph("Open risks", styles["Heading2"]),
                  tbl([["Risk", "Score", "Status"]] + [[Paragraph(esc(r[0]), cell), r[1], esc(r[2])] for r in risks],
                      [400, 60, 100])]
    doc.build(story)
    return buf.getvalue()
